"""Deterministic document classification and inert, bounded previews.

Original bytes never run in the browser. PDF pages are rendered server-side;
workbooks expose cached cell values, not formulas, macros or external links.
"""
from __future__ import annotations

from datetime import date
import hashlib
import io
import os
from pathlib import Path
import re
import socket
import ipaddress
from urllib.parse import urlparse
import zipfile

from db import APP_DATA_DIR
from .store import ControlError

MAX_BYTES = 25 * 1024 * 1024
MIMES = {"PDF": "application/pdf", "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "HTML": "text/html"}


def detect_format(data):
    if data.startswith(b"%PDF-"):
        return "PDF"
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024 or len(z.infolist()) > 10000:
                    raise ControlError("ARCHIVE_TOO_LARGE", "The expanded document exceeds the preview limit.", 422)
                if "xl/workbook.xml" in z.namelist():
                    return "XLSX"
        except zipfile.BadZipFile:
            pass
    if re.search(br"<(?:!doctype\s+html|html|table)\b", data[:4096], re.I):
        return "HTML"
    raise ControlError("UNSUPPORTED_FORMAT", "Only verified PDF, XLSX and HTML documents can be previewed.", 422)


def classify(metadata, header="", detected=None):
    item = dict(metadata)
    blockers, warnings = [], []
    title = header or str(item.get("title") or "")
    # Header evidence takes priority over an external link label.
    interim = re.search(r"(?:six\s+months\s+ended|шесть\s+месяцев|полугод\w*)[^\n]{0,80}?(20\d{2})", title, re.I)
    if interim:
        year = interim.group(1)
        item.update(period_start=f"{year}-01-01", period_end=f"{year}-06-30", duration_months=6,
                    statement_type="interim", period=f"{year}H1", classification_evidence=interim.group(0))
    if detected:
        if item.get("declared_format") and item["declared_format"].upper() != detected:
            warnings.append("FORMAT_MISMATCH")
        item.update(detected_format=detected, detected_mime=MIMES[detected])
    try:
        end = date.fromisoformat(str(item["period_end"])[:10])
        published = date.fromisoformat(str(item.get("published_at") or date.today())[:10])
        if end > published:
            blockers.append("ANNUAL_BEFORE_YEAR_END" if item.get("duration_months") == 12 else "IMPOSSIBLE_PERIOD")
        if end > date.today():
            blockers.append("IMPOSSIBLE_PERIOD")
    except (ValueError, KeyError):
        blockers.append("PERIOD_NOT_CLASSIFIED")
    if not item.get("standard"):
        blockers.append("STANDARD_NOT_CLASSIFIED")
    item.update(blockers=sorted(set(blockers)), warnings=warnings)
    return item


def validate_source_url(url):
    parsed = urlparse(url)
    allowed = {s.strip().lower() for s in os.getenv("ADMIN_DOCUMENT_HOSTS", "openinfo.uz,api.openinfo.uz,uzse.uz").split(",") if s.strip()}
    if parsed.scheme != "https" or parsed.hostname not in allowed or parsed.port not in (None, 443) or parsed.username or parsed.password:
        raise ControlError("SOURCE_NOT_ALLOWED", "This document host is not an approved HTTPS source.", 422)
    try:
        addresses = socket.getaddrinfo(parsed.hostname, 443)
        if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
            raise ControlError("SOURCE_NOT_ALLOWED", "The source resolves to a non-public address.", 422)
    except socket.gaierror:
        raise ControlError("SOURCE_UNAVAILABLE", "The source hostname could not be resolved.", 503) from None


def fetch_original(url):
    import requests
    session = requests.Session()
    session.trust_env = False
    try:
        for _ in range(4):
            validate_source_url(url)
            with session.get(url, stream=True, timeout=(10, 45), allow_redirects=False) as response:
                if response.is_redirect:
                    from urllib.parse import urljoin
                    url = urljoin(url, response.headers.get("Location", ""))
                    continue
                response.raise_for_status()
                data = bytearray()
                for chunk in response.iter_content(65536):
                    data.extend(chunk)
                    if len(data) > MAX_BYTES:
                        raise ControlError("DOCUMENT_TOO_LARGE", "The document exceeds 25 MB.", 422)
                detect_format(bytes(data))
                return bytes(data)
        raise ControlError("TOO_MANY_REDIRECTS", "The document source redirected too many times.", 422)
    finally:
        session.close()


def original_path(checksum):
    if not re.fullmatch(r"[a-f0-9]{64}", checksum or ""):
        raise ControlError("ORIGINAL_UNAVAILABLE", "A verified original has not been stored.", 404)
    return Path(os.getenv("ADMIN_DOCUMENT_DIR") or APP_DATA_DIR / "admin_documents") / checksum


def save_original(data):
    checksum = hashlib.sha256(data).hexdigest()
    path = original_path(checksum)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        # Atomic replace; a crashed download cannot leave a partial original.
        import uuid
        temporary = path.with_suffix("." + uuid.uuid4().hex + ".tmp")
        temporary.write_bytes(data)
        temporary.replace(path)
    return checksum


def read_original(document):
    path = original_path(document.get("checksum"))
    if not path.is_file():
        raise ControlError("ORIGINAL_UNAVAILABLE", "Reprocess the document to store a verified original.", 404)
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != document["checksum"]:
        raise ControlError("CHECKSUM_MISMATCH", "The stored original failed its integrity check.")
    return data


def preview(document, *, page=1, sheet=None, start=0, query=""):
    data = read_original(document)
    kind = detect_format(data)
    if kind == "XLSX":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            name = sheet or book.sheetnames[0]
            if name not in book.sheetnames:
                raise ControlError("SHEET_NOT_FOUND", "The requested worksheet does not exist.", 404)
            ws = book[name]
            rows = []
            scanned = 0
            for row_number, cells in enumerate(ws.iter_rows(min_row=max(1, start + 1), max_row=min(ws.max_row or 2000, 2000), max_col=min(ws.max_column or 100, 100)), start + 1):
                scanned = row_number
                from openpyxl.utils import get_column_letter
                values = [{"address": f"{get_column_letter(column)}{row_number}", "value": str(cell.value) if cell.value is not None else None} for column, cell in enumerate(cells, 1)]
                if not query or query.lower() in " ".join(str(c["value"] or "") for c in values).lower():
                    rows.append(values)
                if len(rows) >= 100:
                    break
            return {"format": kind, "sheets": book.sheetnames, "sheet": name, "rows": rows,
                    "next_row": scanned if scanned < min(ws.max_row or 0, 2000) else None,
                    "row_limit": 2000, "column_limit": 100, "formula_policy": "cached_values_only"}
        finally:
            book.close()
    if kind == "PDF":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(data)
        try:
            if page < 1 or page > len(pdf):
                raise ControlError("PAGE_NOT_FOUND", "The requested page does not exist.", 404)
            pdf_page = pdf[page - 1]
            textpage = pdf_page.get_textpage()
            text = textpage.get_text_range()
            textpage.close()
            pdf_page.close()
            return {"format": kind, "page": page, "pages": len(pdf), "text": text,
                    "matches": [m.start() for m in re.finditer(re.escape(query), text, re.I)] if query else []}
        finally:
            pdf.close()
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "iframe", "object"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)[:200000]
    return {"format": kind, "text": text, "matches": [m.start() for m in re.finditer(re.escape(query), text, re.I)] if query else []}


def pdf_page_image(document, page):
    import pypdfium2 as pdfium
    data = read_original(document)
    if detect_format(data) != "PDF":
        raise ControlError("NOT_PDF", "This document is not a PDF.", 422)
    pdf = pdfium.PdfDocument(data)
    try:
        if not 1 <= page <= len(pdf):
            raise ControlError("PAGE_NOT_FOUND", "The requested page does not exist.", 404)
        p = pdf[page - 1]
        width, height = p.get_size()
        bitmap = p.render(scale=min(1.5, 2000 / max(width, height)))
        output = io.BytesIO()
        bitmap.to_pil().save(output, format="PNG")
        bitmap.close()
        p.close()
        return output.getvalue()
    finally:
        pdf.close()
