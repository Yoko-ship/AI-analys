"""Characterize report content before moving it out of analysis orchestration.

The fixture records the pre-refactor results, including period selection,
numeric cells, translations, and narrative. Export checks inspect real files.
"""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pdfplumber
import pytest

import reporting.exports as subject
from reporting.article.builder import _build_article_report


def report_input():
    def filing(year, assets, cash, equity, income):
        rows = [
            ("Дата отчетности", [f"{year}-12-31"]),
            ("Активы", []),
            ("Касса и платёжные документы", [cash / 2, cash]),
            ("ИТОГО АКТИВОВ", [assets / 2, assets]),
            ("Обязательства и собственный капитал", []),
            ("ИТОГО ОБЯЗАТЕЛЬСТВ", [(assets - equity) / 2, assets - equity]),
            ("ИТОГО СОБСТВЕННОГО КАПИТАЛА", [equity / 2, equity]),
            ("Отчет о финансовых результатах", []),
            ("ИТОГО ПРОЦЕНТНЫХ ДОХОДОВ", [income * 5, income * 10]),
            ("ИТОГО ПРОЦЕНТНЫХ РАСХОДОВ", [income, income * 2]),
            ("ЧИСТАЯ ПРИБЫЛЬ", [income / 2, income]),
        ]
        return {
            "report_id": f"fixture-{year}", "report_form": "NSBU",
            "period_type": "annual", "published_at": f"{year + 1}-03-01",
            "sheets": [{"sheet": "Balance", "table_rows": [
                {"row": i + 1, "label": label, "values": values,
                 "numeric_values": values if values and isinstance(values[0], (int, float)) else []}
                for i, (label, values) in enumerate(rows)
            ]}],
        }

    return {
        "company_name": "Fixture Bank", "ticker": "FIX",
        "annual_period": "2024", "quarterly_period": "",
        "sections": {"ДОСЬЕ": "Public reporting fixture.", "ИТОГ": "Reported results."},
        "metrics": {},
        "ifrs_snapshot": {
            "bank": {"is_bank": True},
            "balance_sheet": {"total_assets": 1000, "equity": 200},
            "income_statement": {"revenue": 300, "net_income": 30},
        },
        "company_data": {"excel_reports": {"items": [
            filing(2023, 800, 80, 160, 20),
            filing(2024, 1000, 100, 200, 30),
        ]}},
    }


@pytest.mark.parametrize("language", ["ru", "en", "uz"])
@pytest.mark.parametrize("mode", ["latest", "annual", "empty"])
def test_article_content_and_input_are_preserved(language, mode):
    kwargs = report_input()
    if mode == "annual":
        kwargs["report_comparison"] = {
            "mode": "annual", "current_year": 2024, "previous_year": 2023,
            "report_form": "NSBU",
        }
    elif mode == "empty":
        kwargs["company_data"] = None
    original = deepcopy(kwargs)
    expected = json.loads(Path(__file__).with_name("fixtures").joinpath("article_reports.json").read_text(encoding="utf-8"))
    result = _build_article_report(**kwargs, language=language)
    assert result == expected[f"{mode}-{language}"]
    assert kwargs == original, "Rendering must not modify collected filings"
    if mode == "annual":
        selection = result["meta"]["report_comparison"]
        assert selection["current_report"]["year"] == 2024
        assert selection["previous_report"]["year"] == 2023
        assert result["meta"]["table_count"] >= 5


def test_missing_requested_period_is_rejected():
    with pytest.raises(ValueError, match="2022"):
        _build_article_report(**report_input(), language="en", report_comparison={
            "mode": "annual", "current_year": 2024, "previous_year": 2022,
            "report_form": "NSBU",
        })


@pytest.mark.parametrize("language,info,metrics", [
    ("en", "Info", "Metrics"), ("ru", "Информация", "Показатели"),
    ("uz", "Ma'lumot", "Ko'rsatkichlar"),
])
def test_excel_preserves_numeric_cells_dates_and_disclaimer(language, info, metrics):
    result = {
        "company_name": "Fixture & Bank", "model": "fixture-model",
        "metrics": {"total_score": {"score": 0, "grade": "C"},
                    "industry": {"debt_equity": {"value": "1,234.5x"}}},
        "sections": {"Overview": "Reported results."},
    }
    at = datetime(2025, 2, 1, 10, 30)
    wb = load_workbook(BytesIO(subject.build_analysis_excel(result, language, at)))
    assert wb[info]["B3"].value == "Fixture & Bank"
    assert wb[info]["B4"].value == at
    cells = dict(wb[metrics].iter_rows(min_row=2, values_only=True))
    assert cells["Score / Итоговый скор"] == 0
    assert cells["D/E"] == 1234.5
    assert cells["Grade / Класс"] == "C"
    assert subject.report_disclaimer(language) in [cell.value for row in wb[info] for cell in row]


@pytest.mark.parametrize("language", ["en", "ru", "uz"])
def test_pdf_contains_company_narrative_and_disclaimer(language):
    payload = subject.build_analysis_pdf({
        "company_name": "Fixture & Bank", "ticker": "FIX",
        "sections": {"Overview": "Reported results."},
    }, language, datetime(2025, 2, 1))
    with pdfplumber.open(BytesIO(payload)) as pdf:
        text = " ".join(page.extract_text() for page in pdf.pages)
    normalized = " ".join(text.split())
    assert "Fixture & Bank" in text
    assert "Reported results." in text
    assert " ".join(subject.report_disclaimer(language).split()) in normalized


def test_comparison_exports_share_the_same_matrix():
    result = {"comparison": {
        "fields": [{"key": "score", "label": "Score"}],
        "rows": [{"company_name": "First", "score": 0},
                 {"company_name": "Second", "score": None}],
    }}
    at = datetime(2025, 2, 1)
    wb = load_workbook(BytesIO(subject.build_comparison_excel(result, "en", at)))
    assert list(wb["Metrics"].values) == [("Metric", "First", "Second"), ("Score", 0, "—")]
    payload = subject.build_comparison_pdf(result, "en", at)
    with pdfplumber.open(BytesIO(payload)) as pdf:
        text = " ".join(page.extract_text() for page in pdf.pages)
    assert all(value in text for value in ("First", "Second", "Score"))
