import json
from pathlib import Path

from fastapi.testclient import TestClient
import pytest

from fastapi import FastAPI
from server.company.financial_routes import router
from catalogue.storage import get_catalog_conn
from financial_columns import source_store

app = FastAPI()
app.include_router(router)


@pytest.mark.parametrize('statement,form', [('balance','BANK_BS'),('income','BANK_PL')])
@pytest.mark.parametrize('view', ['template','overview'])
def test_statement_tabs_and_exports_keep_only_the_requested_statement(source, statement, form, view):
    import csv
    import io
    from openpyxl import load_workbook
    query=f'view={view}&statement={statement}'
    with TestClient(app) as client:
        data=client.get(f'/api/company/IPTB/financials/columns?{query}').json()
        csv_result=client.get(f'/api/company/IPTB/financials/columns.csv?{query}')
        xlsx_result=client.get(f'/api/company/IPTB/financials/columns.xlsx?{query}')
    assert data['rows'] and {r['source_form'] for r in data['rows']} == {form}
    records=list(csv.DictReader(io.StringIO(csv_result.content.decode('utf-8-sig'))))
    assert {r['source_form'] for r in records} == {form}
    worksheet=list(load_workbook(io.BytesIO(xlsx_result.content)).active.values)
    assert {r[worksheet[0].index('source_form')] for r in worksheet[1:]} == {form}
    assert any(r['value'] not in {'',None} for r in records)


def test_statement_filter_rejects_unknown_names(source):
    with TestClient(app) as client:
        assert client.get('/api/company/IPTB/financials/columns?statement=BAD').status_code == 422


@pytest.fixture
def source(monkeypatch, tmp_path):
    monkeypatch.setenv('CATALOG_DB_PATH', str(tmp_path / 'catalog.db'))
    c = get_catalog_conn()
    with c:
        c.execute("INSERT INTO catalog_companies(ticker,company_name,org_id) VALUES ('IPTB','Ipoteka bank','217')")
    c.close()
    data = json.loads(Path('tests/fixtures/openinfo/iptb_2026q3_complete_rows.json').read_text(encoding='utf-8'))
    source_store.record_nsbu(ticker='IPTB', issuer_id='217', year=2026, quarter=3, template='bank',
                            data={'ok': True, 'income': data, 'balance': data},
                            urls={'excel_url': data['source_url'], 'published_at': '2026-10-20'})
    return data


def test_public_matrix_uses_full_stored_forms_and_keeps_stocks_out_of_ttm(source):
    with TestClient(app) as client:
        response = client.get('/api/company/IPTB/financials/columns?form=NSBU&scope=separate&freq=interim')
    assert response.status_code == 200
    result = response.json()
    assert result['unit'] == 'UZS_thousand'
    income = next(r for r in result['rows'] if r['id'] == 'BANK_PL:1л')
    cell = income['cells']['2026-09-30']
    assert cell['value'] == '9447836814'
    assert cell['source_url'] == source['source_url']
    assert cell['source_code'] == '1л'
    assert cell['scope'] == 'separate'
    assert cell['raw_value'] == 9447836814
    assert result['source_forms']
    assert result['coverage']['notes'] == 'not_applicable'


@pytest.mark.parametrize('view', ['template', 'overview'])
def test_iptb_annual_screen_and_export_include_current_report_amounts(source, view):
    import csv
    import io
    with TestClient(app) as client:
        response = client.get(f'/api/company/IPTB/financials/columns?view={view}')
        exported = client.get(f'/api/company/IPTB/financials/columns.csv?view={view}')
    assert response.status_code == exported.status_code == 200
    result = response.json()
    assert [(c['label'], c['basis']) for c in result['columns']] == [('2025', 'annual'), ('2026 · 9M', 'ytd')]
    current = result['columns'][-1]['key']
    income = next(r for r in result['rows'] if r['source_form'] == 'BANK_PL' and r['source_code'] == '1л')
    assert income['cells'][current]['value'] == '9447836814'
    assert income['cells'][current]['period_start'] == '2026-01-01'
    assert income['cells'][current]['source_url'] == source['source_url']
    rows = list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
    assert any(r['source_code'] == '1л' and r['value'] == '9447836814' and r['period_end'] == '2026-09-30' for r in rows)


def test_matrix_rejects_invalid_context_and_never_reads_nsbu_as_consolidated(source):
    with TestClient(app) as client:
        assert client.get('/api/company/IPTB/financials/columns?form=BAD').status_code == 422
        assert client.get('/api/company/IPTB/financials/columns?freq=BAD').status_code == 422
        result = client.get('/api/company/IPTB/financials/columns?scope=consolidated').json()
    assert result['rows'] == []


def test_csv_and_excel_export_include_source_context_and_zero_cells(source):
    import csv
    import io
    from openpyxl import load_workbook
    with TestClient(app) as client:
        csv_response = client.get('/api/company/IPTB/financials/columns.csv?freq=interim')
        excel_response = client.get('/api/company/IPTB/financials/columns.xlsx?freq=interim')
    assert csv_response.status_code == excel_response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
    item = next(r for r in rows if r['source_code'] == '1л' and r['period_end'] == '2026-09-30')
    assert item['standard'] == 'NSBU' and item['scope'] == 'separate'
    assert item['source_url'] == source['source_url']
    assert item['quality'] == 'valid'
    assert item['value'] == '9447836814'
    assert any(r['value'] == '0' for r in rows)
    workbook = load_workbook(io.BytesIO(excel_response.content))
    assert list(workbook.active.values)[0] == tuple(rows[0].keys())
    assert len(list(workbook.active.values)) == len(rows) + 1


def test_openinfo_ratios_are_separate_from_uzstock_calculations(source):
    from catalogue.facts import upsert_facts
    upsert_facts([{'entity_id': '217', 'dataset': 'financial_indicators', 'field': 'roe',
                   'period': '2025', 'value': 12.5, 'unit': '%',
                   'source': 'openinfo_financial_indicators', 'source_url': source['source_url']},
                  {'entity_id': '217', 'dataset': 'financial_indicators', 'field': 'roe',
                   'period': '2025', 'value': 99, 'unit': '%',
                   'source': 'nsbu_derived', 'source_url': source['source_url']}])
    with TestClient(app) as client:
        result = client.get('/api/company/IPTB/financials/columns').json()
    assert len(result['published_ratios']) == 1
    assert result['published_ratios'][0]['value'] == '12.5'
    assert result['published_ratios'][0]['origin'] == 'OpenInfo_published'
    assert all(cell['origin'] == 'UZStock_calculated' for r in result['ratios'] for cell in r['cells'].values())


def test_exports_include_published_ratios_owners_and_literal_spreadsheet_labels():
    import csv
    import io
    from financial_columns.exports import export
    payload = {'ticker':'UZNF', 'issuer_id':'1059', 'standard':'MSFO', 'scope':'separate',
               'rows':[], 'ratios':[], 'columns':[],
               'published_ratios':[{'field':'roe', 'period':'2024', 'value':'-12.5', 'unit':'%',
                                    'source_url':'https://openinfo.uz', 'origin':'OpenInfo_published'}],
               'shareholders':[{'name':'-HYPERLINK("https://example.com")', 'share':'75', 'date':'2024-12-31',
                                'quality':'valid', 'share_class':'ordinary', 'page':6, 'source_url':'https://openinfo.uz'}]}
    content, _ = export(payload, 'csv')
    rows = list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
    assert next(r for r in rows if r['metric_id'] == 'roe')['value'] == '-12.5'
    owner = next(r for r in rows if r['source_form'] == 'OWNERS')
    assert owner['label'].startswith("'-HYPERLINK")
    assert owner['share_class'] == 'ordinary'
    assert owner['page'] == '6'


def test_real_leasing_form_exposes_calculated_debt_and_the_leasing_ratio(source):
    c = get_catalog_conn()
    with c:
        c.execute("INSERT INTO catalog_companies(ticker,company_name,org_id) VALUES ('UZAL','O‘zagrolizing','638')")
    c.close()
    data = json.loads(Path('tests/fixtures/openinfo/uzal_2025_complete_rows.json').read_text(encoding='utf-8'))
    source_store.record_nsbu(ticker='UZAL', issuer_id='638', year=2025, quarter=0, template='financial',
                            data={'ok':True, 'balance':data, 'income':data}, urls={'excel_url':data.get('source_url', data.get('url'))})
    with TestClient(app) as client:
        result = client.get('/api/company/UZAL/financials/columns').json()
    assert result['template'] == 'leasing'
    debt = next(r for r in result['rows'] if r['metric_id'] == 'interest_bearing_debt')
    assert debt['cells']['2025']['value'] is not None
    assert len(debt['cells']['2025']['components']) == 5
    cost = next(r for r in result['ratios'] if r['metric_id'] == 'funding_cost')
    assert cost['cells']['2025']['quality'] == 'valid'
    assert float(cost['cells']['2025']['value']) > 0


def test_reference_view_keeps_only_named_bank_fields_and_exports_the_same_rows(source):
    import csv
    import io
    with TestClient(app) as client:
        raw = client.get('/api/company/IPTB/financials/columns?freq=interim').json()
        selected = client.get('/api/company/IPTB/financials/columns?freq=interim&view=template').json()
        exported = client.get('/api/company/IPTB/financials/columns.csv?freq=interim&view=template')
    assert selected['presentation'] == 'template'
    codes = {(r['source_form'], r['source_code']) for r in selected['rows']}
    assert ('BANK_PL', '1л') in codes
    assert ('BANK_PL', '1а') in codes
    assert ('BANK_PL', '7з') in codes
    assert ('BANK_BS', '14') in codes
    assert ('BANK_PL', '7а') not in codes  # staff expense is not a reference field
    assert ('BANK_PL', '7в') not in codes  # travel expense
    assert ('BANK_PL', '1') not in codes   # heading must not duplicate the total
    assert ('BANK_BS', '31') not in codes # unlisted balance check total
    assert len(selected['rows']) < len(raw['rows'])
    for row in selected['rows']:
        assert row['cells'] == next(r for r in raw['rows'] if r['id'] == row['id'])['cells']
    income = next(r for r in selected['rows'] if r['id'] == 'BANK_PL:1л')
    raw_income = next(r for r in raw['rows'] if r['id'] == income['id'])
    assert income['label'] == 'Итого процентных доходов'
    assert income['source_label'] == raw_income['label']
    assert income['cells']['2026-09-30']['raw_label'] == raw_income['label']
    rows = list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
    assert next(r for r in rows if r['source_code'] == '1л')['label'] == income['label']
    assert not any((r['source_form'], r['source_code']) in {('BANK_PL','7а'), ('BANK_PL','7в'), ('BANK_BS','31')} for r in rows)
    assert any(r['source_code'] == '1л' and r['value'] == '9447836814' for r in rows)
    assert all(r['metric_id'] in {'roe','roa','nim','cir','cost_of_risk','loans_to_deposits'} for r in selected['ratios'])


def test_reference_view_rejects_unknown_view(source):
    with TestClient(app) as client:
        assert client.get('/api/company/IPTB/financials/columns?view=bad').status_code == 422
        assert client.get('/api/company/IPTB/financials/columns.csv?view=bad').status_code == 422


def test_overview_export_matches_the_compact_html_list(source):
    import csv
    import io
    with TestClient(app) as client:
        response = client.get('/api/company/IPTB/financials/columns?freq=interim&view=overview')
        exported = client.get('/api/company/IPTB/financials/columns.csv?freq=interim&view=overview')
    assert response.status_code == exported.status_code == 200
    data = response.json()
    assert len(data['rows']) == 24
    rows = list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
    assert {r['label'] for r in rows} == {r['label'] for r in data['rows']}
    assert not any(r['source_code'] == '7з' and r['source_form'] == 'BANK_PL' for r in rows)
    assert any(r['source_code'] == '7а' and r['source_form'] == 'BANK_BS' and r['value'] for r in rows)
