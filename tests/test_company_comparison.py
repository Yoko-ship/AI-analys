"""Real filing cases for the deterministic company comparison contract."""
from copy import deepcopy
from datetime import date
import importlib

import pytest
from fastapi.testclient import TestClient

# These are reported NSBU values, in thousand UZS, already covered by the
# catalog's real-case tests. SQBN FY2024 is its known line-number/zero filing.
HMKB_2025 = dict(revenue=5_174_244_859, net_income=1_704_969_118,
                 total_assets=35_911_649_727, total_equity=6_704_709_904,
                 total_liabilities=29_206_939_823, gross_profit=3_956_102_734,
                 operating_expenses=1_890_152_285)
HMKB_2024 = dict(revenue=4_353_666_010, net_income=1_426_604_717,
                 total_assets=28_708_226_712, total_equity=5_161_403_031,
                 total_liabilities=23_546_823_681, gross_profit=3_168_483_582,
                 operating_expenses=1_413_924_153)
HMKB_Q3 = dict(revenue=4_645_444_343, net_income=1_480_422_343,
               total_assets=44_389_278_550, total_equity=8_185_132_247,
               total_liabilities=36_204_146_303, gross_profit=3_577_446_147,
               operating_expenses=1_819_867_178)
SQBN_BAD = dict(revenue=1, net_income=13, total_assets=0,
                total_equity=0, total_liabilities=0, gross_profit=14, operating_expenses=6)


@pytest.fixture
def source(monkeypatch):
    service = importlib.import_module('server.comparison.service')
    import catalogue.snapshots as snapshots
    import catalogue.ratios as ratios
    import catalogue.filings as filings
    import catalogue.organization as org
    import financial_ingestion.publication as publication
    from reporting import store
    records = {
        'HMKB': {'ticker': 'HMKB', 'name': 'Hamkorbank', 'type': 'stock', 'isin': 'UZHMKB', 'sector': 'finance'},
        'HMKBP': {'ticker': 'HMKBP', 'name': 'Hamkorbank', 'type': 'stock', 'isin': 'UZHMKBP', 'sector': 'finance', 'is_preferred': True},
        'SQBN': {'ticker': 'SQBN', 'name': 'SQB bank', 'type': 'stock', 'isin': 'UZSQBN', 'sector': 'finance'},
        'UZTL': {'ticker': 'UZTL', 'name': 'Uztelecom', 'type': 'stock', 'isin': 'UZUZTL', 'sector': 'telecom'},
        'BOND': {'ticker': 'BOND', 'name': 'Bond', 'type': 'bond'},
    }
    annual = {'HMKB': {'2025': HMKB_2025, '2024': HMKB_2024}, 'SQBN': {'2024': SQBN_BAD}, 'UZTL': {}}
    quarters = {'HMKB': {'2026Q3': HMKB_Q3}, 'SQBN': {}, 'UZTL': {}}
    monkeypatch.setattr(service, 'get_securities_map', lambda: deepcopy(records))
    monkeypatch.setattr(snapshots, 'get_financials_series', lambda t, form='NSBU': deepcopy(annual.get(t.replace('HMKBP', 'HMKB'), {})))
    monkeypatch.setattr(snapshots, 'get_financials_series_quarterly', lambda t, form='NSBU': deepcopy(quarters.get(t.replace('HMKBP', 'HMKB'), {})))
    monkeypatch.setattr(ratios, '_filed_periods', lambda t, form: {'2025', '2024', '2026Q3'})
    monkeypatch.setattr(org, 'filing_org_types', lambda ts: {t: 'bank' for t in ts if t != 'UZTL'})
    def reports(t):
        return [{'report_form': 'NSBU', 'year': int(p[:4]), 'quarter': int(p[-1]) if 'Q' in p else 0,
                 'pdf_url': 'https://openinfo.uz/ru/reports/to_pdf31438/', 'published_at': '2026-10-19', 'report_id': 31438}
                for p in [*annual.get(t, {}), *quarters.get(t, {})]]
    monkeypatch.setattr(filings, 'get_company_reports', reports)
    monkeypatch.setattr(filings, 'get_company_index', lambda t: {'org_id': t, 'tickers': ['HMKB', 'HMKBP'] if t in ['HMKB', 'HMKBP'] else [t]})
    monkeypatch.setattr(publication, 'snapshots', lambda t, scope=None: [])
    monkeypatch.setattr(store, 'latest_publication', lambda *args: None)
    monkeypatch.setattr(service, 'market_rows', lambda securities: {
        'HMKB': {**records['HMKB'], 'last_price': 100, 'last_trade_date': date.today().isoformat(), 'shares_outstanding': 1_000_000_000},
        'HMKBP': {**records['HMKBP'], 'last_price': 1000, 'last_trade_date': date.today().isoformat(), 'shares_outstanding': 10_000_000},
    })
    monkeypatch.setattr(service.market_store, 'get_all_trade_stats', lambda: {})
    return service, annual, quarters


def test_real_hamkorbank_statement_is_scaled_once_and_source_linked(source):
    service, _, _ = source
    result = service.compare(['HMKB'], '2025', 'NSBU', 'separate')
    col = result['companies'][0]
    assert result['period'] == '2025'
    assert result['period_basis'] == 'annual'
    assert col['metrics']['net_profit']['value'] == HMKB_2025['net_income'] * 1000
    assert col['metrics']['interest_income']['value'] == HMKB_2025['revenue'] * 1000
    assert col['metrics']['revenue']['status'] == 'not_applicable'
    assert col['metrics']['net_profit']['source']['url'].startswith('https://openinfo.uz/')
    assert col['metrics']['net_profit']['change'] == pytest.approx((HMKB_2025['net_income'] / HMKB_2024['net_income'] - 1) * 100)
    assert col['metrics']['roe']['value'] == pytest.approx(HMKB_2025['net_income'] / ((HMKB_2024['total_equity'] + HMKB_2025['total_equity']) / 2) * 100)
    assert col['metrics']['roe']['formula']
    assert 'ranking' not in result and 'summary' not in result


def test_cumulative_nine_month_figures_are_not_discrete_quarters(source):
    service, _, _ = source
    result = service.compare(['HMKB'], '2026Q3', 'NSBU', 'separate')
    assert result['period_basis'] == 'cumulative_ytd'
    metrics = result['companies'][0]['metrics']
    assert metrics['net_profit']['value'] == 1_480_422_343_000
    assert metrics['roe']['value'] == pytest.approx(HMKB_Q3['net_income'] * 12 / 9 / ((HMKB_2025['total_equity'] + HMKB_Q3['total_equity']) / 2) * 100)
    assert metrics['roe']['annualised'] is True
    assert metrics['net_profit']['change'] is None


def test_exact_period_never_falls_back_and_junk_filings_never_become_figures(source):
    service, _, _ = source
    result = service.compare(['HMKB', 'SQBN'], '2024', 'NSBU', 'separate')
    assert result['companies'][0]['metrics']['net_profit']['value'] == HMKB_2024['net_income'] * 1000
    assert result['companies'][1]['metrics']['net_profit']['value'] is None
    assert result['companies'][1]['status'] == 'unavailable'
    missing = service.compare(['SQBN'], '2025', 'NSBU', 'separate')
    assert missing['companies'][0]['metrics']['assets']['value'] is None


@pytest.mark.parametrize('standard,scope', [('NSBU', 'consolidated'), ('IFRS', 'separate'), ('IFRS', 'consolidated')])
def test_standards_and_scopes_never_borrow_nsbu_figures(source, standard, scope):
    service, _, _ = source
    result = service.compare(['HMKB'], '2025', standard, scope)
    assert result['companies'][0]['metrics']['net_profit']['value'] is None


def test_latest_uses_one_shared_period_not_each_companys_latest(source):
    service, annual, _ = source
    annual['SQBN']['2025'] = {**HMKB_2025, 'net_income': 1_880_069_362}
    result = service.compare(['HMKB', 'SQBN'], 'latest', 'NSBU', 'separate')
    assert result['period'] == '2025'
    assert all(c['metrics']['net_profit']['period'] == '2025' for c in result['companies'])


def test_issuer_capitalisation_covers_both_share_classes(source):
    service, _, _ = source
    result = service.compare(['HMKB', 'HMKBP'], '2025', 'NSBU', 'separate')
    left, right = result['companies']
    assert left['metrics']['market_cap']['value'] == 110_000_000_000
    assert right['metrics']['market_cap']['value'] == left['metrics']['market_cap']['value']
    assert left['metrics']['pe']['value'] == right['metrics']['pe']['value']
    assert left['metrics']['pb']['value'] == pytest.approx(110_000_000_000 / (HMKB_2025['total_equity'] * 1000))
    assert left['metrics']['pb']['value'] == right['metrics']['pb']['value']
    assert left['metrics']['price']['value'] != right['metrics']['price']['value']


def test_options_include_real_universe_and_dynamic_periods(source):
    service, annual, _ = source
    annual['HMKB']['2023'] = HMKB_2024
    result = service.options(['HMKB'], 'NSBU', 'separate')
    assert '2023' in result['periods']
    assert 'BOND' not in [c['ticker'] for c in result['companies']]
    assert set(c['ticker'] for c in result['companies']) >= {'HMKB', 'SQBN', 'UZTL'}


def test_public_http_comparison_validates_and_has_v2_alias(source):
    import api
    client = TestClient(api.app)
    for path in ['/api/company-comparison', '/api/v2/company-comparison']:
        response = client.get(path, params={'tickers': 'HMKB,SQBN', 'period': '2025'})
        assert response.status_code == 200
        assert response.json()['companies'][0]['metrics']['net_profit']['value'] == HMKB_2025['net_income'] * 1000
    for params in [{'tickers': 'HMKB,SQBN,UZTL,HMKBP,BOND'}, {'tickers': '../secret'}, {'tickers': 'HMKB', 'period': '9999Q9'}, {'tickers': 'HMKB', 'scope': 'anything'}]:
        assert client.get('/api/company-comparison', params=params).status_code == 422
    assert client.get('/api/company-comparison', params={'tickers': 'UNKNOWN'}).status_code == 404


def test_excel_is_generated_from_server_values_with_sources(source):
    import api
    from io import BytesIO
    from openpyxl import load_workbook
    response = TestClient(api.app).get('/api/company-comparison/export.xlsx?tickers=HMKB&period=2025')
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert set(book.sheetnames) == {'Comparison', 'Context', 'Sources'}
    rows = list(book['Comparison'].values)
    assert any(row[0] == 'net_profit' and row[1] == 1_704_969_118_000 for row in rows)
    assert any('https://openinfo.uz/' in str(value) for row in book['Sources'].values for value in row)


def test_published_ifrs_scope_keeps_its_own_units_sources_and_bank_fields(source, monkeypatch):
    service, _, _ = source
    import financial_ingestion.publication as publication
    payload = {'classification': {'standard': 'MSFO', 'scope': 'consolidated',
               'period_start': '2025-01-01', 'period_end': '2025-12-31', 'document_year': 2025},
               'normalized_uzs': {'net_income': '1704969118000', 'total_assets': '35911649727000',
               'total_equity': '6704709904000', 'total_liabilities': '29206939823000',
               'loans_to_customers': '20000000000000', 'customer_deposits': '17000000000000'},
               'source': {'pdf_url': 'https://openinfo.uz/ifrs-source.pdf', 'file_hash': 'a' * 64, 'published_at': '2026-05-15'}}
    monkeypatch.setattr(publication, 'snapshots', lambda t, scope=None: [{'id': 'published-2025', 'created_at': '2026-10-20', 'payload': payload}] if scope == 'consolidated' else [])
    col = service.compare(['HMKB'], '2025', 'IFRS', 'consolidated')['companies'][0]
    assert col['metrics']['net_profit']['value'] == 1704969118000
    assert col['source']['url'] == 'https://openinfo.uz/ifrs-source.pdf'
    assert col['source']['publication_date'] == '2026-05-15'
    assert col['metrics']['loan_book']['value'] == 20000000000000
    assert col['metrics']['customer_funds']['value'] == 17000000000000
    assert service.compare(['HMKB'], '2025', 'IFRS', 'separate')['companies'][0]['metrics']['net_profit']['value'] is None


def test_verified_bank_lines_only_supplement_the_same_period_and_scope(source, monkeypatch):
    service, _, _ = source
    from reporting import store
    report = {'status': 'available', 'period': '2026Q3', 'scope': 'separate', 'verified_facts': [
        {'metric_code': 'loan_portfolio', 'value': 30_908_613_408, 'unit': 'thousand UZS',
         'verification_status': 'verified', 'period': '2026Q3', 'consolidation_scope': 'separate',
         'accounting_standard': 'NSBU', 'source_url': 'https://openinfo.uz/ru/reports/to_pdf31438/'}]}
    monkeypatch.setattr(store, 'latest_publication', lambda *a: report)
    assert service.compare(['HMKB'], '2026Q3')['companies'][0]['metrics']['loan_book']['value'] == 30_908_613_408_000
    assert service.compare(['HMKB'], '2025')['companies'][0]['metrics']['loan_book']['value'] is None
    report['scope'] = 'consolidated'
    assert service.compare(['HMKB'], '2026Q3')['companies'][0]['metrics']['loan_book']['value'] is None


def test_nsbu_separate_filing_is_never_presented_as_a_consolidated_source(source):
    service, _, _ = source
    col = service.compare(['HMKB'], '2025', 'NSBU', 'consolidated')['companies'][0]
    assert col['source'] is None
    assert col['metrics']['net_profit']['source'] is None
